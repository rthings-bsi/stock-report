#!/usr/bin/env python3
"""
Parser & Aggregator Produktivitas Gudang SAP MB51.
Mendukung:
1. File Mentah / Raw Export SAP MB51 (28 kolom: Entry Date, Entry Time, Plant, Sloc, Posting Date,
   Mvt, Name 2, Purchase Order, Order, Work center, Kode Material, Material Description, Batch,
   Qty, Berat, Rp., Document Header Text, Material Document, Doc. Item, Reference, GR/GI Slip,
   User name, Text, Unloading Point, Sales Order, Sales order item, KG GI, KG GR)
2. File Excel Daily Report WH (Sheet 'MB51' dengan 46 kolom lengkap termasuk rumus)
3. Format .xlsx, .xls (termasuk TSV/text export dari SAP GUI), .csv, .tsv, .txt

Seluruh rumus klasifikasi (Jenis Matr., Gudang, HandHeld/Desktop, Tanggal Kerja, Shift, Jam,
Masuk/Keluar, Kode Work Center, dan Kode Work Center + MVT) dihitung otomatis 100% identik
dengan sheet 'ah' dan 'MB51' pada buku kerja Daily Report WH.
"""
import argparse
import csv
import datetime
import io
import json
import os
import sys
from collections import defaultdict

GUDANG_ORDER = [
    'Gd.01', 'Gd.02', 'Gd.03', 'Gd.04', 'Gd.05', 'Gd.06',
    'Gd.08', 'Gd.09', 'Gd.10', 'Gd.11', 'Gd.12', 'Gd.13', 'Gd.14',
    'x Sloc Lama', 'x Packaging', 'x Transit'
]

# Tabel VLOOKUP(LEFT(Kode Material, 2), ah!D:E, 2, FALSE)
KODMAT_PREFIX_MAP = {
    'C1': '1. Coil',
    'C2': '2. Strip',
    'C3': '2. Strip',
    'A2': '3. Pipa', 'A4': '3. Pipa', 'G1': '3. Pipa', 'G5': '3. Pipa',
    'H1': '3. Pipa', 'H2': '3. Pipa', 'H3': '3. Pipa', 'J1': '3. Pipa',
    'P3': '3. Pipa', 'U1': '3. Pipa', 'UA': '3. Pipa', 'UB': '3. Pipa',
    'X1': '3. Pipa', 'X3': '3. Pipa', 'X4': '3. Pipa', 'XA': '3. Pipa',
    'XB': '3. Pipa', 'XC': '3. Pipa', 'XD': '3. Pipa', 'XF': '3. Pipa',
    'XG': '3. Pipa', 'XH': '3. Pipa', 'XJ': '3. Pipa', 'Y1': '3. Pipa',
    'Y3': '3. Pipa', 'YA': '3. Pipa', 'YB': '3. Pipa', 'YC': '3. Pipa',
    'YD': '3. Pipa', 'YE': '3. Pipa', 'YG': '3. Pipa', 'YH': '3. Pipa',
    'Z1': '3. Pipa', 'Z5': '3. Pipa', 'ZA': '3. Pipa', 'ZC': '3. Pipa',
}

# Tabel VLOOKUP(MID(Sloc, 2, 1), ah!J:K, 2, FALSE)
SLOC_TO_GUDANG = {
    'A': 'Gd.01',
    'B': 'Gd.02',
    'C': 'Gd.03',
    'D': 'Gd.04',
    'E': 'Gd.05',
    'F': 'Gd.06',
    'G': 'Gd.08',
    'H': 'Gd.09',
    'I': '-',
    'J': 'Gd.10',
    'K': 'Gd.11',
    'L': 'Gd.12',
    'M': 'Gd.13',
    'N': 'Gd.14',
    'O': 'x Packaging',
    'T': 'x Transit',
    '5': 'x Sloc Lama',
}

# Tabel VLOOKUP(User name, ah!R:S, 2, FALSE)
USER_TO_DEVICE = {
    'ADMWHBARCODE': 'Desktop',
    'ADMWHEXISTNG': 'Desktop',
    'ADMWHEXPANSI': 'Desktop',
    'ADMWHRTP': 'Desktop',
    'AHADI': 'Desktop',
    'ARONI': 'Desktop',
    'DFIRDAUS': 'Desktop',
    'DTARIGAN': 'Desktop',
    'ERINALDY': 'Desktop',
    'FARDIAWAN': 'Desktop',
    'FPERDANA': 'Desktop',
    'GIGR-03': 'Desktop',
    'GIPURWAKA': 'Desktop',
    'GUDANG01A': 'HandHeld',
    'GUDANG01B': 'HandHeld',
    'GUDANG02A': 'HandHeld',
    'GUDANG02B': 'HandHeld',
    'GUDANG02C': 'HandHeld',
    'GUDANG02D': 'HandHeld',
    'GUDANG03A': 'HandHeld',
    'GUDANG03B': 'HandHeld',
    'GUDANG04B': 'HandHeld',
    'GUDANG04C': 'HandHeld',
    'GUDANG05A': 'HandHeld',
    'GUDANG05B': 'HandHeld',
    'GUDANG05C': 'HandHeld',
    'GUDANG05D': 'HandHeld',
    'GUDANG05E': 'HandHeld',
    'GUDANG10A': 'HandHeld',
    'GUDANG11A': 'HandHeld',
    'GUDANG11B': 'HandHeld',
    'GUDANG11C': 'HandHeld',
    'GUDANG12A': 'HandHeld',
    'GUDANG12B': 'HandHeld',
    'GUDANG13A': 'HandHeld',
    'GUDANG13B': 'HandHeld',
    'GUDANG13C': 'HandHeld',
    'GUDANG14A': 'HandHeld',
    'GUDANG14B': 'HandHeld',
    'GUDANG14C': 'HandHeld',
    'GUDANGADM01': 'Desktop',
    'GUDANGADM02': 'Desktop',
    'GUDANGADM03': 'Desktop',
    'GUDANGADM04': 'Desktop',
    'GUDANGADM05': 'Desktop',
    'GUDANGADM08': 'Desktop',
    'GUDANGRM1': 'HandHeld',
    'GUDANGRM2': 'HandHeld',
    'GUDANGRM3': 'HandHeld',
    'GUDANGRTP01': 'HandHeld',
    'GUDANGRTP02': 'HandHeld',
    'HSAPUTRA': 'Desktop',
    'IHAKIM': 'Desktop',
    'IRAFIANSYAH': 'Desktop',
    'ISETYAWAN': 'Desktop',
    'OMARYANTO': 'Desktop',
    'QC_SKF-001': 'Desktop',
    'RHAMID': 'Desktop',
    'RMAULIDA': 'Desktop',
    'RPURBA': 'Desktop',
    'RRUD-01': 'Desktop',
    'RRUD-02': 'Desktop',
    'RRUD-04': 'Desktop',
    'SRACHIM': 'Desktop',
    'TSEPTIADI': 'Desktop',
}

# Tabel VLOOKUP(Mvt & " - " & IF(Qty<0,"OUT","IN"), ah!U:V, 2, FALSE)
MVT_DIR_MAP = {
    ('101', 'IN'): '1. Penerimaan',
    ('102', 'OUT'): '1. Penerimaan',
    ('261', 'OUT'): '2. Pengeluaran',
    ('262', 'IN'): '2. Pengeluaran',
    ('309', 'IN'): '-',
    ('309', 'OUT'): '-',
    ('310', 'IN'): '-',
    ('310', 'OUT'): '-',
    ('311', 'IN'): '1. Penerimaan',
    ('311', 'OUT'): '2. Pengeluaran',
    ('312', 'IN'): '1. Penerimaan',
    ('312', 'OUT'): '2. Pengeluaran',
    ('601', 'OUT'): '2. Pengeluaran',
    ('602', 'IN'): '2. Pengeluaran',
    ('641', 'IN'): '2. Pengeluaran',
    ('641', 'OUT'): '2. Pengeluaran',
    ('642', 'OUT'): '2. Pengeluaran',
    ('901', 'IN'): '1. Penerimaan',
    ('902', 'OUT'): '1. Penerimaan',
    ('541', 'OUT'): '2. Pengeluaran',
    ('542', 'IN'): '2. Pengeluaran',
}

# Tabel VLOOKUP(Kode Work Center, ah!X:Y, 2, FALSE)
WC_LOOKUP_MAP = {
    '-': '-',
    '- 4800': '0. SPINDO Unit Lain',
    '- 4900': '0. SPINDO Unit Lain',
    '- 5000': 'no',
    'ALUMINIUM EXTRUSION INDONESIA - 4500': '6. Subcont Alexindo',
    'BAOSTEEL SINGAPORE PTE LTD. - 4700': '0. Vendor Import',
    'CUTC - 5000': '3.b Cut. CIRCLE',
    'CUTH - 5000': '3.c Cut. HELLES',
    'CUTL - 5000': '3.d Cut. LASER',
    'CUTP - 5000': '3.a Cut. HAVEN',
    'CUTS - 5000': '3.e Cut. SAW',
    'CV. CAHAYA SINAR GEMILANG - 4500': '6. Subcont CSG',
    'DEB- - 5000': '4.c DEBURING',
    'EBR- - 5000': '4.b ENDBRUSHING',
    'EFC- - 5000': 'no',
    'EXSW - 5000': '4.a SWADGING',
    'HYD- - 5000': 'no',
    'M&P- - 5000': '5. Marking Packing',
    'MILL - 5000': '2. Prod. MILL',
    'PT. RAJAWALI KARYA PERKASA - 4500': '6. Subcont RKP',
    'PT. VASTEX INTI MULIA - 4500': '6. Subcont VASTEX',
    'PT. WELDPART INDONESIA - 4500': '6. Subcont WPI',
    'PT.BAOSTEEL INDONESIA STEEL - 4500': '6. Subcont BISC',
    'REP- - 5000': 'no',
    'SINO GLORY METAL RESOURCES INTERNAT - 4700': '0. Vendor Import',
    'SLIT - 5000': '1. Prod. SLITTER',
    'UTF- - 5000': 'no',
    'POSCO INTERNATIONAL CORPORATION - 4700': '0. Vendor Import',
    'PT. LAUTAN BAJA INDONESIA - 4500': '0. Vendor LOCAL',
    'PT. PDF PRESISI ENGINEERING - 4500': '6. Subcont PDF Presisi',
    'PT UNIVERSAL METAL INVESTMENT - 4500': '0. Vendor LOCAL',
}

# Tabel VLOOKUP(Kode Work Center & " / " & Mvt, ah!AA:AB, 2, FALSE)
WC_MVT_LOOKUP_MAP = {
    '- / 309': 'no',
    '- / 310': 'no',
    '- / 311': '5. Transfer SLOC',
    '- / 312': '5. Transfer SLOC',
    '- / 601': '4. Kirim Customer',
    '- / 602': '4. Kirim Customer',
    '0. SPINDO UNIT LAIN / 101': '0. SPINDO Unit Lain',
    '0. SPINDO UNIT LAIN / 641': '0. SPINDO Unit Lain',
    '0. SPINDO UNIT LAIN / 642': '0. SPINDO Unit Lain',
    '0. VENDOR IMPORT / 901': '3. Vendor',
    '1. PROD. SLITTER / 101': '1. Hasil Produksi',
    '1. PROD. SLITTER / 102': '1. Hasil Produksi',
    '1. PROD. SLITTER / 261': '2. Bahan Produksi',
    '1. PROD. SLITTER / 262': '2. Bahan Produksi',
    '2. PROD. MILL / 101': '1. Hasil Produksi',
    '2. PROD. MILL / 102': '1. Hasil Produksi',
    '2. PROD. MILL / 261': '2. Bahan Produksi',
    '2. PROD. MILL / 262': '2. Bahan Produksi',
    '3.A CUT. HAVEN / 101': '1. Hasil Produksi',
    '3.A CUT. HAVEN / 102': '1. Hasil Produksi',
    '3.A CUT. HAVEN / 261': '2. Bahan Produksi',
    '3.A CUT. HAVEN / 262': '2. Bahan Produksi',
    '3.B CUT. CIRCLE / 101': '1. Hasil Produksi',
    '3.B CUT. CIRCLE / 102': '1. Hasil Produksi',
    '3.B CUT. CIRCLE / 261': '2. Bahan Produksi',
    '3.B CUT. CIRCLE / 262': '2. Bahan Produksi',
    '3.C CUT. HELLES / 101': '1. Hasil Produksi',
    '3.C CUT. HELLES / 102': '1. Hasil Produksi',
    '3.C CUT. HELLES / 261': '2. Bahan Produksi',
    '3.C CUT. HELLES / 262': '2. Bahan Produksi',
    '3.D CUT. LASER / 101': '1. Hasil Produksi',
    '3.D CUT. LASER / 102': '1. Hasil Produksi',
    '3.D CUT. LASER / 261': '2. Bahan Produksi',
    '3.D CUT. LASER / 262': '2. Bahan Produksi',
    '3.E CUT. SAW / 101': '1. Hasil Produksi',
    '3.E CUT. SAW / 102': '1. Hasil Produksi',
    '3.E CUT. SAW / 261': '2. Bahan Produksi',
    '3.E CUT. SAW / 262': '2. Bahan Produksi',
    '4.A SWADGING / 101': '1. Hasil Produksi',
    '4.A SWADGING / 102': '1. Hasil Produksi',
    '4.A SWADGING / 261': '2. Bahan Produksi',
    '4.A SWADGING / 262': '2. Bahan Produksi',
    '4.B ENDBRUSHING / 101': '1. Hasil Produksi',
    '4.B ENDBRUSHING / 261': '2. Bahan Produksi',
    '4.B ENDBRUSHING / 262': '2. Bahan Produksi',
    '4.C DEBURING / 101': '1. Hasil Produksi',
    '4.C DEBURING / 261': '2. Bahan Produksi',
    '4.C DEBURING / 262': '2. Bahan Produksi',
    '5. MARKING PACKING / 101': '1. Hasil Produksi',
    '5. MARKING PACKING / 102': '1. Hasil Produksi',
    '5. MARKING PACKING / 261': '2. Bahan Produksi',
    '5. MARKING PACKING / 262': '2. Bahan Produksi',
    '6. SUBCONT ALEXINDO / 901': '3. Subcont',
    '6. SUBCONT BISC / 901': '3. Subcont',
    '6. SUBCONT BISC / 902': '3. Subcont',
    '6. SUBCONT CSG / 901': '3. Subcont',
    '6. SUBCONT CSG / 902': '3. Subcont',
    '6. SUBCONT RKP / 901': '3. Subcont',
    '6. SUBCONT RKP / 902': '3. Subcont',
    '6. SUBCONT VASTEX / 901': '3. Subcont',
    '6. SUBCONT VASTEX / 902': '3. Subcont',
    '6. SUBCONT WPI / 901': '3. Subcont',
    '6. SUBCONT WPI / 902': '3. Subcont',
    '6. SUBCONT PDF PRESISI / 901': '3. Subcont',
    '0. VENDOR LOCAL / 901': '3. Vendor',
    '0. VENDOR LOCAL / 902': '3. Vendor',
    'NO / 101': 'no',
    'NO / 102': 'no',
    'NO / 261': 'no',
    'NO / 262': 'no',
    '- / 541': '4. Kirim SUBCONT',
    '- / 542': '4. Kirim SUBCONT',
}


def norm_code_str(val):
    """Normalisasi nilai sel (kode Mvt, PO, Order, Sloc, Username) tanpa akhiran .0 dari float Excel."""
    if val is None:
        return ''
    if isinstance(val, int):
        return str(val)
    if isinstance(val, float):
        return str(int(val)) if val.is_integer() else str(val)
    s = str(val).strip()
    if not s or s == 'None' or s.startswith('#'):
        return ''
    if s.endswith('.0') and s[:-2].lstrip('-').isdigit():
        return s[:-2]
    return s


def format_date(val):
    """Mendukung datetime, YYYY-MM-DD, DD.MM.YYYY, DD/MM/YYYY, dan DD:MM:YYYY (format export SAP MB51)."""
    if val is None:
        return ''
    if isinstance(val, datetime.datetime):
        return val.strftime('%Y-%m-%d')
    if isinstance(val, datetime.date):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    if not s or s.startswith('#'):
        return ''
    s10 = s[:10].strip()
    if len(s10) == 10 and s10[4] == '-' and s10[7] == '-':
        return s10
    for fmt in ('%d:%m:%Y', '%d.%m.%Y', '%d/%m/%Y', '%Y/%m/%d', '%Y:%m:%d', '%m/%d/%Y'):
        try:
            return datetime.datetime.strptime(s10, fmt).strftime('%Y-%m-%d')
        except ValueError:
            pass
    return ''


def parse_hour(val):
    if val is None:
        return None
    if isinstance(val, datetime.time):
        return val.hour
    if isinstance(val, datetime.datetime):
        return val.hour
    if isinstance(val, datetime.timedelta):
        total_seconds = int(val.total_seconds()) % 86400
        return total_seconds // 3600
    if isinstance(val, (int, float)):
        if 0.0 <= float(val) < 1.0:
            return int(round(float(val) * 86400) // 3600) % 24
        if 0 <= val <= 23 and float(val).is_integer():
            return int(val)
    s = str(val).strip()
    if ':' in s:
        try:
            return int(s.split(':')[0])
        except ValueError:
            return None
    if '.' in s:
        try:
            h = int(s.split('.')[0])
            if 0 <= h <= 23:
                return h
        except ValueError:
            return None
    return None


def parse_float(val):
    """Mendukung angka float/int, spasi setelah tanda minus ('-   184'), minus di belakang ('184-'), dan pemisah ribuan."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s or s.startswith('#') or s == '-':
        return 0.0

    is_neg = False
    if s.startswith('-') or s.endswith('-'):
        is_neg = True
        s = s.strip('-').strip()
    elif s.startswith('(') and s.endswith(')'):
        is_neg = True
        s = s[1:-1].strip()

    s = s.replace(' ', '').replace(' ', '')
    if not s:
        return 0.0

    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        parts = s.split(',')
        if len(parts) == 2 and len(parts[1]) <= 3:
            s = s.replace(',', '.')
        else:
            s = s.replace(',', '')
    try:
        res = float(s)
        return -res if is_neg else res
    except ValueError:
        return 0.0


def derive_jenis_matr(kode_material_val):
    """=VLOOKUP(LEFT(AC6,2),ah!D:E,2,FALSE)"""
    km = norm_code_str(kode_material_val).upper()
    if len(km) >= 2:
        prefix = km[:2]
        if prefix in KODMAT_PREFIX_MAP:
            return KODMAT_PREFIX_MAP[prefix]
    return '3. Pipa'


def derive_gudang(sloc_val):
    """=VLOOKUP(MID(V6,2,1),ah!J:K,2,FALSE)"""
    s = norm_code_str(sloc_val).upper()
    if len(s) >= 2:
        return SLOC_TO_GUDANG.get(s[1], '')
    return ''


def derive_device(username_val):
    """=VLOOKUP(AN6,ah!R:S,2,FALSE) dengan fallback prefix GUDANG (non-GUDANGADM) untuk akun handheld baru."""
    u = norm_code_str(username_val).upper()
    if not u:
        return ''
    if u in USER_TO_DEVICE:
        return USER_TO_DEVICE[u]
    if u.startswith('GUDANG') and not u.startswith('GUDANGADM'):
        return 'HandHeld'
    return ''


def derive_shift_and_tk(entry_date_str, hour):
    """
    Shift: =VLOOKUP(T6,ah!M:N,2)
    Tanggal Kerja: =IF(VLOOKUP(T6,ah!M:P,4)="Day-1",S6-1,S6)
    """
    if hour is None:
        return 'Shift 1', entry_date_str
    if 7 <= hour <= 14:
        return 'Shift 1', entry_date_str
    if 15 <= hour <= 22:
        return 'Shift 2', entry_date_str
    if 0 <= hour <= 6 and entry_date_str:
        try:
            dt = datetime.datetime.strptime(entry_date_str, '%Y-%m-%d') - datetime.timedelta(days=1)
            return 'Shift 3', dt.strftime('%Y-%m-%d')
        except ValueError:
            pass
    return 'Shift 3', entry_date_str


def derive_masuk_keluar(mvt_val, qty_val):
    """=VLOOKUP(X6&" - "&IF(AF6<0,"OUT","IN"),ah!U:V,2,FALSE)"""
    mvt_code = norm_code_str(mvt_val)
    dir_flag = 'OUT' if qty_val < 0 else 'IN'
    return MVT_DIR_MAP.get((mvt_code, dir_flag), '')


def derive_wc_and_mvt(name2_val, po_val, order_val, wc_val, mvt_val):
    """
    Col N (Kode Work Center):
      =IFERROR(VLOOKUP(IF(Z6="","",Y6)&LEFT(AB6,4)&" - "&LEFT(AA6,4)&LEFT(Z6,4),ah!X:Y,2,FALSE),
               IF(Z6="","",Y6)&LEFT(AB6,4)&" - "&LEFT(AA6,4)&LEFT(Z6,4))
    Col O (Kode Work Center + MVT):
      =IFERROR(VLOOKUP(N6&" / "&X6,ah!AA:AB,2,FALSE),"no data")
    """
    name2 = norm_code_str(name2_val)
    po = norm_code_str(po_val)
    order = norm_code_str(order_val)
    wc = norm_code_str(wc_val)
    mvt = norm_code_str(mvt_val)

    left_part = (name2 if po else '') + wc[:4]
    right_part = order[:4] + po[:4]
    wc_key = f'{left_part} - {right_part}'.strip().upper()
    kode_wc = WC_LOOKUP_MAP.get(wc_key, wc_key)
    wc_mvt_key = f'{kode_wc} / {mvt}'.strip().upper()
    kode_wc_mvt = WC_MVT_LOOKUP_MAP.get(wc_mvt_key, 'no data')
    return kode_wc, kode_wc_mvt


def find_header_map(row_values):
    norm = {}
    for idx, val in enumerate(row_values):
        if val is not None:
            key = ' '.join(str(val).strip().lower().split())
            if key and key not in norm:
                norm[key] = idx
    return norm


def get_col_idx(norm_map, candidates):
    for c in candidates:
        k = ' '.join(c.strip().lower().split())
        if k in norm_map:
            return norm_map[k]
    return None


def process_rows(rows_iter, source_filename, last_updated_hint=''):
    header_map = None
    col_indices = {}

    buckets = defaultdict(lambda: {
        'hh_in_kg': 0.0,
        'hh_out_kg': 0.0,
        'hh_in_tx': 0,
        'hh_out_tx': 0,
        'all_in_kg': 0.0,
        'all_out_kg': 0.0,
        'all_in_tx': 0,
        'all_out_tx': 0,
        'raw_abs_kg': 0.0,
        'raw_tx': 0,
        'coil_kg': 0.0,
        'strip_kg': 0.0,
        'pipa_kg': 0.0,
    })

    hourly_buckets = defaultdict(lambda: {'hh_tx': 0, 'all_tx': 0, 'hh_kg': 0.0})

    total_rows_read = 0
    dates_pd_set = set()
    dates_tk_set = set()

    for row_idx, row in enumerate(rows_iter):
        if header_map is None:
            if row_idx > 30:
                break
            cand_map = find_header_map(row)
            has_date = get_col_idx(
                cand_map,
                ['entry date', 'posting date', 'pstng date', 'tanggal kerja', 'entered on', 'cpudt', 'budat']
            ) is not None
            has_berat_or_qty = get_col_idx(
                cand_map,
                ['berat', 'weight', 'qty', 'quantity', 'quantity in une']
            ) is not None
            if has_date and has_berat_or_qty:
                header_map = cand_map
                col_indices = {
                    # Kolom helper Excel (opsional, tidak ada pada file raw MB51)
                    'jenis_matr': get_col_idx(header_map, ['jenis matr.', 'jenis material', 'jenis mtr.2']),
                    'gudang': get_col_idx(header_map, ['gudang', 'warehouse']),
                    'device': get_col_idx(header_map, ['handheld / desktop', 'hh / desktop', 'device']),
                    'tanggal_kerja': get_col_idx(header_map, ['tanggal kerja']),
                    'shift': get_col_idx(header_map, ['shift']),
                    'jam': get_col_idx(header_map, ['jam', 'jam kerja']),
                    'masuk_keluar': get_col_idx(header_map, ['masuk / keluar', 'in / out']),
                    'kode_wc': get_col_idx(header_map, ['kode work center']),
                    'kode_wc_mvt': get_col_idx(header_map, ['kode work center + mvt']),
                    # 28 Kolom Raw SAP MB51
                    'entry_date': get_col_idx(header_map, ['entry date', 'entered on', 'cpudt']),
                    'entry_time': get_col_idx(header_map, ['entry time', 'time of entry', 'entered at', 'cputm']),
                    'plant': get_col_idx(header_map, ['plant', 'werks']),
                    'sloc': get_col_idx(header_map, ['sloc', 'storage location', 'stor. loc.', 'lgort']),
                    'posting_date': get_col_idx(header_map, ['posting date', 'pstng date', 'budat']),
                    'mvt': get_col_idx(header_map, ['mvt', 'movement type', 'mov.type', 'bwart']),
                    'name2': get_col_idx(header_map, ['name 2', 'name 1', 'vendor', 'customer']),
                    'po': get_col_idx(header_map, ['purchase order', 'po', 'ebeln']),
                    'order': get_col_idx(header_map, ['order', 'aufnr']),
                    'work_center': get_col_idx(header_map, ['work center', 'work ctr', 'arbpl']),
                    'kode_material': get_col_idx(header_map, ['kode material', 'material', 'matnr']),
                    'material_desc': get_col_idx(header_map, ['material description', 'material desc', 'maktx']),
                    'batch': get_col_idx(header_map, ['batch', 'charg']),
                    'qty': get_col_idx(
                        header_map,
                        ['qty in un. of entry', 'qty in un.of entry', 'quantity in une', 'qty in unit of entry', 'qty', 'erfmg']
                    ),
                    'berat': get_col_idx(
                        header_map,
                        ['berat', 'weight', 'net weight', 'ntgew', 'brgew', 'quantity', 'menge']
                    ),
                    'username': get_col_idx(header_map, ['user name', 'username', 'created by', 'usnam']),
                }
                if col_indices['berat'] is None:
                    col_indices['berat'] = col_indices['qty']
            continue

        if not row:
            continue

        def val(key):
            idx = col_indices.get(key)
            if idx is not None and idx < len(row):
                return row[idx]
            return None

        raw_berat = val('berat')
        raw_tk = val('tanggal_kerja')
        raw_pd = val('posting_date')
        raw_ed = val('entry_date')

        if raw_berat is None and raw_tk is None and raw_pd is None and raw_ed is None:
            continue

        berat = parse_float(raw_berat)
        qty_val = parse_float(val('qty'))

        ed_str = format_date(raw_ed)
        pd_str = format_date(raw_pd)
        tk_str = format_date(raw_tk)

        # Jika Posting Date di-mask '########' pada export Excel, gunakan Entry Date
        if not pd_str and ed_str:
            pd_str = ed_str
        if not ed_str and pd_str:
            ed_str = pd_str

        hour = parse_hour(val('jam'))
        if hour is None:
            hour = parse_hour(val('entry_time'))

        sh_raw = val('shift')
        sh = norm_code_str(sh_raw)
        if not sh or not tk_str:
            derived_sh, derived_tk = derive_shift_and_tk(ed_str, hour)
            if not sh:
                sh = derived_sh
            if not tk_str:
                tk_str = derived_tk

        sh_upper = sh.upper()
        if '1' in sh_upper:
            sh = 'Shift 1'
        elif '2' in sh_upper:
            sh = 'Shift 2'
        elif '3' in sh_upper:
            sh = 'Shift 3'
        else:
            continue

        if not tk_str and pd_str:
            tk_str = pd_str
        if not pd_str and tk_str:
            pd_str = tk_str
        if not tk_str and not pd_str:
            continue

        total_rows_read += 1

        # 1. Gudang (dari kolom Gudang atau karakter ke-2 Sloc)
        gd = norm_code_str(val('gudang'))
        if not gd:
            gd = derive_gudang(val('sloc'))
        if not gd or gd == '-':
            continue

        # 2. HandHeld / Desktop (dari kolom Device atau VLOOKUP User name)
        dev = norm_code_str(val('device'))
        if not dev:
            dev = derive_device(val('username'))
        is_handheld = dev.lower() == 'handheld'

        # 3. Masuk / Keluar (dari kolom Masuk/Keluar atau VLOOKUP Mvt + tanda Qty)
        mk = norm_code_str(val('masuk_keluar'))
        if not mk:
            sign_ref = qty_val if qty_val != 0 else berat
            mk = derive_masuk_keluar(val('mvt'), sign_ref)

        # 4. Kode Work Center + MVT (dari kolom O atau rumus Name 2 + PO + Order + Work center + Mvt)
        wc_mvt = norm_code_str(val('kode_wc_mvt'))
        if not wc_mvt:
            _, wc_mvt = derive_wc_and_mvt(
                val('name2'),
                val('po'),
                val('order'),
                val('work_center'),
                val('mvt'),
            )
        is_mvt_valid = wc_mvt.lower() != 'no'

        # 5. Jenis Material (dari kolom Jenis Matr. atau 2 karakter awal Kode Material)
        jm = norm_code_str(val('jenis_matr'))
        if not jm:
            jm = derive_jenis_matr(val('kode_material'))

        if pd_str:
            dates_pd_set.add(pd_str)
        if tk_str:
            dates_tk_set.add(tk_str)

        for d_mode, d_str in (('posting_date', pd_str), ('tanggal_kerja', tk_str)):
            if not d_str:
                continue
            b = buckets[(d_mode, d_str, sh, gd)]
            b['raw_abs_kg'] += abs(berat)
            b['raw_tx'] += 1

            if is_mvt_valid:
                if mk == '1. Penerimaan':
                    b['all_in_kg'] += berat
                    b['all_in_tx'] += 1
                    if is_handheld:
                        b['hh_in_kg'] += berat
                        b['hh_in_tx'] += 1
                elif mk == '2. Pengeluaran':
                    b['all_out_kg'] += berat
                    b['all_out_tx'] += 1
                    if is_handheld:
                        b['hh_out_kg'] += berat
                        b['hh_out_tx'] += 1

                if is_handheld and mk in ('1. Penerimaan', '2. Pengeluaran'):
                    abs_b = abs(berat)
                    if 'coil' in jm.lower():
                        b['coil_kg'] += abs_b
                    elif 'strip' in jm.lower():
                        b['strip_kg'] += abs_b
                    else:
                        b['pipa_kg'] += abs_b

            if hour is not None and is_mvt_valid and mk in ('1. Penerimaan', '2. Pengeluaran'):
                hb = hourly_buckets[(d_mode, d_str, sh, hour, gd)]
                hb['all_tx'] += 1
                if is_handheld:
                    hb['hh_tx'] += 1
                    hb['hh_kg'] += abs(berat)

    records = []
    for (d_mode, d_str, sh, gd), b in sorted(buckets.items()):
        records.append({
            'dateMode': d_mode,
            'tanggal': d_str,
            'shift': sh,
            'gudang': gd,
            'hhInKg': round(b['hh_in_kg'], 3),
            'hhOutKg': round(b['hh_out_kg'], 3),
            'hhInTx': b['hh_in_tx'],
            'hhOutTx': b['hh_out_tx'],
            'allInKg': round(b['all_in_kg'], 3),
            'allOutKg': round(b['all_out_kg'], 3),
            'allInTx': b['all_in_tx'],
            'allOutTx': b['all_out_tx'],
            'rawAbsKg': round(b['raw_abs_kg'], 3),
            'rawTx': b['raw_tx'],
            'coilKg': round(b['coil_kg'], 1),
            'stripKg': round(b['strip_kg'], 1),
            'pipaKg': round(b['pipa_kg'], 1),
        })

    hourly_records = []
    for (d_mode, d_str, sh, hour, gd), hb in sorted(hourly_buckets.items()):
        if hb['hh_tx'] > 0 or hb['all_tx'] > 0:
            hourly_records.append({
                'dateMode': d_mode,
                'tanggal': d_str,
                'shift': sh,
                'hour': hour,
                'gudang': gd,
                'hhTx': hb['hh_tx'],
                'allTx': hb['all_tx'],
                'hhKg': round(hb['hh_kg'], 1),
            })

    return {
        'sourceFile': os.path.basename(source_filename),
        'lastUpdated': last_updated_hint or datetime.datetime.now().strftime('%d.%m.%Y - %H:%M WIB'),
        'totalRowsRead': total_rows_read,
        'postingDates': sorted(dates_pd_set),
        'tanggalKerjaDates': sorted(dates_tk_set),
        'records': records,
        'hourlyRecords': hourly_records,
    }


def parse_delimited_text(filepath):
    """Membaca file CSV / TSV / Text export dari SAP GUI (mendukung UTF-8, UTF-16, CP1252)."""
    raw_bytes = b''
    with open(filepath, 'rb') as bf:
        raw_bytes = bf.read()

    text = ''
    for enc in ('utf-8-sig', 'utf-16', 'cp1252', 'latin-1'):
        try:
            text = raw_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        text = raw_bytes.decode('utf-8', errors='replace')

    sample = text[:8192]
    for sep in ('\t', ';', ',', '|'):
        if sample.count(sep) > 5:
            reader = csv.reader(io.StringIO(text), delimiter=sep)
            res = process_rows(reader, filepath)
            if res['totalRowsRead'] > 0:
                return res

    reader = csv.reader(io.StringIO(text))
    return process_rows(reader, filepath)


def parse_file(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext in ('.csv', '.tsv', '.txt'):
        return parse_delimited_text(filepath)

    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        target_sheet = 'MB51' if 'MB51' in wb.sheetnames else None
        if not target_sheet:
            for sn in wb.sheetnames:
                if 'mb51' in sn.lower():
                    target_sheet = sn
                    break

        last_update_hint = ''
        for sn in ('Produktivitas Harian', 'PivoT1', 'PivoT2'):
            if sn in wb.sheetnames:
                ws_meta = wb[sn]
                for r in ws_meta.iter_rows(min_row=1, max_row=25, max_col=5, values_only=True):
                    for cell in r:
                        if cell and 'last update' in str(cell).lower():
                            last_update_hint = (
                                str(cell)
                                .replace('Last Update :', '')
                                .replace('Last Update MB51 :', '')
                                .strip()
                            )
                            break
                    if last_update_hint:
                        break
            if last_update_hint:
                break

        sheets_to_try = [target_sheet] if target_sheet else list(wb.sheetnames)
        for sn in wb.sheetnames:
            if sn not in sheets_to_try:
                sheets_to_try.append(sn)

        for sn in sheets_to_try:
            ws = wb[sn]
            res = process_rows(ws.iter_rows(values_only=True), filepath, last_update_hint)
            if res['totalRowsRead'] > 0:
                wb.close()
                return res

        wb.close()
    except Exception:
        # Fallback jika file .xls adalah text/TSV export langsung dari SAP GUI
        return parse_delimited_text(filepath)

    return parse_delimited_text(filepath)


def self_check(json_path):
    assert os.path.exists(json_path), f'Missing {json_path}'
    with open(json_path, 'r', encoding='utf-8') as f:
        d = json.load(f)
    assert d['totalRowsRead'] > 0, f"Expected >0 rows, got {d['totalRowsRead']}"
    if d['totalRowsRead'] == 70373:
        assert len(d['postingDates']) == 10
        rec_09 = [
            r for r in d['records']
            if r['dateMode'] == 'posting_date' and r['tanggal'] == '2025-12-09'
            and r['gudang'] == 'Gd.01' and r['shift'] == 'Shift 1'
        ][0]
        hh_09 = rec_09['hhInKg'] + abs(rec_09['hhOutKg'])
        assert abs(hh_09 - 8352.364) < 0.01, f'Expected 8352.364, got {hh_09}'
        rec_10 = [
            r for r in d['records']
            if r['dateMode'] == 'tanggal_kerja' and r['tanggal'] == '2025-12-10'
            and r['gudang'] == 'Gd.06' and r['shift'] == 'Shift 1'
        ][0]
        assert round(rec_10['rawAbsKg']) == 258203 and rec_10['rawTx'] == 137, f"Got {rec_10}"

    # Verify Raw 28-Column SAP MB51 sample 1 (Daily Report headers: Sloc, Mvt, Kode Material, Qty, Berat)
    raw_sample_tsv = "\n".join([
        "Entry Date\tEntry Time\tPlant\tSloc\tPosting Date\tMvt\tName 2\tPurchase Order\tOrder\tWork center\tKode Material\tMaterial Description\tBatch\tQty\t Berat \tRp.\tDocument Header Text\tMaterial Document\tDoc. Item\tReference\tGR/GI Slip\tUser name\tText\tUnloading Point\tSales Order\tSales order item\tKG GI\tKG GR",
        "01:12:2025\t07:02\t1105\t5N21\t########\t601\tPT ASTEMO BEKASI MANUFACTURING\t\t\t\tYBB12BIA0200+04310\tPM 12BIBRB MILL CUT DEB MP 1-1/2\" 38,1x2\t5240914HEA\t-240 \t- 184 \t0 \t\t4910782546\t8\t8010674986\t\tRPURBA\t\t\t2010282594\t4\t0\t0",
        "01:12:2025\t07:02\t1105\t5N21\t01:12:2025\t601\tPT ASTEMO BEKASI MANUFACTURING\t\t\t\tYBB12BKO0160+02710\tPM 12BIBRB MILL CUT DEB MP  42,7x1,6x271\t5232131HGA\t-114 \t-50 \t0 \t\t4910782546\t15\t8010674986\t\tRPURBA\t\t\t2010282594\t7\t0\t0",
        "01:12:2025\t07:02\t1105\t5N21\t01:12:2025\t601\tPT ASTEMO BEKASI MANUFACTURING\t\t\t\tYBB12BKO0160+02710\tPM 12BIBRB MILL CUT DEB MP  42,7x1,6x271\t5243203HJA\t-126 \t-55 \t0 \t\t4910782546\t16\t8010674986\t\tRPURBA\t\t\t2010282594\t7\t0\t0",
        "01:12:2025\t07:02\t1105\t5N21\t01:12:2025\t601\tPT ASTEMO BEKASI MANUFACTURING\t\t\t\tYBB12BKO0160+02770\tPM 12BIBRB MILL CUT DEB MP  42,7x1,6x277\t5243203HJA\t-120 \t-54 \t0 \t\t4910782546\t3\t8010674986\t\tGUDANG14A\t\t\t2010282594\t8\t0\t0",
    ])
    raw_res = process_rows(csv.reader(io.StringIO(raw_sample_tsv), delimiter='\t'), 'raw_mb51.txt')
    assert raw_res['totalRowsRead'] == 4, f"Expected 4 raw rows, got {raw_res['totalRowsRead']}"
    rec_raw = [r for r in raw_res['records'] if r['dateMode'] == 'posting_date'][0]
    assert rec_raw['tanggal'] == '2025-12-01' and rec_raw['gudang'] == 'Gd.14' and rec_raw['shift'] == 'Shift 1'
    assert rec_raw['rawAbsKg'] == 343.0 and rec_raw['allOutKg'] == -343.0 and rec_raw['hhOutKg'] == -54.0

    # Verify Direct SAP MB51 Export sample 2 (Time of Entry, Storage Location, Movement Type, Material, Qty in Un. of Entry, Quantity)
    sap_direct_tsv = "\n".join([
        "Entry Date\tTime of Entry\tPlant\tStorage Location\tPosting Date\tMovement Type\tName 2\tPurchase Order\tOrder\tWork center\tMaterial\tMaterial Description\tBatch\tQty in Un. of Entry\tQuantity\tAmount in LC\tDocument Header Text\tMaterial Document\tMaterial Doc.Item\tReference\tGoods Receipt/Issue Slip\tUser name\tText\tUnloading Point\tSales Order\tSales order item\tKG GI\tKG GR",
        "24/10/2026\t00:14:56\t1105\t5M09\t24/10/2026\t311\tPT. BERDIKARI METAL ENGINEERING\t\t\t\tXCB12BAA0320+03390\tPM 13ACIBRB MILL CUT DEB MP 1\" 25,4x3,20\t5262490HJA\t-1000\t-594,00\t0,00\t\t4912548907\t1\t\t\tGUDANG13C\t\t4002218823\t2010309195\t3\t0\t0",
        "24/10/2026\t00:14:56\t1105\t5M09\t24/10/2026\t311\tPT. BERDIKARI METAL ENGINEERING\t\t\t\tXCB12BAA0320+03390\tPM 13ACIBRB MILL CUT DEB MP 1\" 25,4x3,20\t5263051HCA\t-224\t-133,056\t0,00\t\t4912548907\t2\t\t\tGUDANG13C\t\t4002218823\t2010309195\t3\t0\t0",
        "24/10/2026\t00:14:56\t1105\t5M09\t24/10/2026\t311\tPT. BERDIKARI METAL ENGINEERING\t\t\t\tXCB12BAA0320+03390\tPM 13ACIBRB MILL CUT DEB MP 1\" 25,4x3,20\t5263052HCA\t-215\t-127,71\t0,00\t\t4912548907\t3\t\t\tGUDANG13C\t\t4002218823\t2010309195\t3\t0\t0",
        "24/10/2026\t00:14:56\t1105\t5M09\t24/10/2026\t311\tPT. BERDIKARI METAL ENGINEERING\t\t\t\tXCB12BAA0320+03390\tPM 13ACIBRB MILL CUT DEB MP 1\" 25,4x3,20\t5263053HCA\t-907\t-538,758\t0,00\t\t4912548907\t4\t\t\tGUDANG13C\t\t4002218823\t2010309195\t3\t0\t0",
        "24/10/2026\t00:14:56\t1105\t5M09\t24/10/2026\t311\tPT. BERDIKARI METAL ENGINEERING\t\t\t\tXCB12BAA0320+03390\tPM 13ACIBRB MILL CUT DEB MP 1\" 25,4x3,20\t5263055HCA\t-93\t-55,242\t0,00\t\t4912548907\t5\t\t\tGUDANG13C\t\t4002218823\t2010309195\t3\t0\t0",
    ])
    sap_res = process_rows(csv.reader(io.StringIO(sap_direct_tsv), delimiter='\t'), 'NYOBA.XLSX')
    assert sap_res['totalRowsRead'] == 5, f"Expected 5 SAP rows, got {sap_res['totalRowsRead']}"
    rec_tk = [r for r in sap_res['records'] if r['dateMode'] == 'tanggal_kerja'][0]
    assert rec_tk['tanggal'] == '2026-10-23' and rec_tk['gudang'] == 'Gd.13' and rec_tk['shift'] == 'Shift 3'
    assert abs(rec_tk['hhOutKg'] - (-1448.766)) < 0.001 and rec_tk['hhOutTx'] == 5 and abs(rec_tk['pipaKg'] - 1448.8) < 0.1

    print('SELF-CHECK OK:', d['totalRowsRead'], 'rows,', len(d['records']), 'buckets + both 28-col MB51 formats verified')


def main():
    parser = argparse.ArgumentParser(description='Parse SAP MB51 file into aggregated JSON')
    parser.add_argument('--input', help='Path to input XLSX/CSV/TSV file')
    parser.add_argument('--output', help='Path to output JSON file')
    parser.add_argument('--self-check', help='Run assert self-check on JSON output file')
    args = parser.parse_args()

    if args.self_check:
        self_check(args.self_check)
        return

    if not args.input or not args.output:
        parser.error('--input and --output are required unless --self-check is used')

    if not os.path.exists(args.input):
        print(f'Error: Input file not found: {args.input}', file=sys.stderr)
        sys.exit(1)

    result = parse_file(args.input)
    if result['totalRowsRead'] == 0:
        print(
            'Error: Tidak ditemukan baris data MB51 yang valid (pastikan file memiliki kolom Entry Date/Posting Date, Sloc, dan Berat/Qty).',
            file=sys.stderr,
        )
        sys.exit(2)

    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
    with open(args.output, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False)

    print(json.dumps({
        'success': True,
        'totalRowsRead': result['totalRowsRead'],
        'postingDatesCount': len(result['postingDates']),
        'recordsCount': len(result['records']),
        'hourlyRecordsCount': len(result['hourlyRecords']),
        'lastUpdated': result['lastUpdated'],
    }))


if __name__ == '__main__':
    main()
